"""Readers: Google Sheet / XLSX / CSV -> list[Table]."""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path

import requests
from openpyxl import load_workbook

from .extract import Cell, Table

_HYPERLINK_FORMULA = re.compile(r'HYPERLINK\(\s*"([^"]+)"\s*(?:[,;]\s*"([^"]*)")?', re.IGNORECASE)
_SHEET_ID = re.compile(r"/spreadsheets/d/([a-zA-Z0-9_-]+)")
_PUBLISHED = re.compile(r"/spreadsheets/d/e/([a-zA-Z0-9_-]+)")


def google_sheet_export_url(url: str) -> str:
    """Whole-workbook XLSX export. XLSX (not CSV) because it keeps every tab
    *and* cell hyperlinks — most tournament sheets show 'MP Link' text with the URL hidden."""
    m = _PUBLISHED.search(url)
    if m:
        return f"https://docs.google.com/spreadsheets/d/e/{m.group(1)}/pub?output=xlsx"
    m = _SHEET_ID.search(url)
    if not m:
        raise ValueError(f"Not a Google Sheets URL: {url}")
    return f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=xlsx"


def read_google_sheet(url: str, timeout: float = 60) -> list[Table]:
    resp = requests.get(google_sheet_export_url(url), timeout=timeout)
    if resp.status_code in (401, 403) or "text/html" in resp.headers.get("content-type", ""):
        raise PermissionError(
            "Google refused the export. Set sharing to 'Anyone with the link can view' "
            "or download the sheet as .xlsx and import the file instead."
        )
    resp.raise_for_status()
    return read_xlsx(io.BytesIO(resp.content))


def read_xlsx(src: Path | io.BytesIO) -> list[Table]:
    # Two passes: formulas (to read =HYPERLINK(...)) and cached values (to read display text).
    if isinstance(src, io.BytesIO):
        raw = src.getvalue()
        wb_f = load_workbook(io.BytesIO(raw), data_only=False)
        wb_v = load_workbook(io.BytesIO(raw), data_only=True)
    else:
        wb_f = load_workbook(src, data_only=False)
        wb_v = load_workbook(src, data_only=True)

    tables: list[Table] = []
    for ws_f in wb_f.worksheets:
        ws_v = wb_v[ws_f.title]
        table = Table(name=ws_f.title)
        for row_f, row_v in zip(ws_f.iter_rows(), ws_v.iter_rows()):
            cells: list[Cell] = []
            for cf, cv in zip(row_f, row_v):
                link = cf.hyperlink.target if getattr(cf, "hyperlink", None) else None
                text = "" if cv.value is None else str(cv.value)
                if isinstance(cf.value, str) and cf.value.startswith("="):
                    m = _HYPERLINK_FORMULA.search(cf.value)
                    if m:
                        link = link or m.group(1)
                        text = text or (m.group(2) or "")
                if text or link:
                    cells.append(Cell(text=text, link=link))
            if cells:
                table.rows.append(cells)
        tables.append(table)
    return tables


def read_csv(path: Path) -> list[Table]:
    text = path.read_text(encoding="utf-8-sig")
    dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    rows = [[Cell(text=v) for v in row if v.strip()] for row in csv.reader(io.StringIO(text), dialect)]
    return [Table(name=path.stem, rows=[r for r in rows if r])]
