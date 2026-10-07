import csv
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from scout import ingest  # noqa: E402
from scout.analytics import build_report  # noqa: E402
from scout.classify import infer_bucket  # noqa: E402
from scout.db import connect, repo  # noqa: E402
from scout.osu.parser import parse_match, parse_title  # noqa: E402
from scout.rounds import normalize_round  # noqa: E402
from scout.sources import discover_links, extract_links_from_text  # noqa: E402
from scout.sources.readers import google_sheet_export_url  # noqa: E402

import fake_tournament as ft  # noqa: E402


# ---------- sources ---------------------------------------------------------
def test_round_normalization():
    assert normalize_round("Grand Finals") == "GF"
    assert normalize_round("Losers Semifinals") == "SF"
    assert normalize_round("RO16 - Day 2") == "RO16"
    assert normalize_round("Quarter-finals") == "QF"
    assert normalize_round("Finals") == "F"
    assert normalize_round("Team Finland") is None
    assert normalize_round("Qualifiers") == "Q"


def test_text_links_and_sections():
    text = """Quarterfinals
    https://osu.ppy.sh/community/matches/111 A vs B
    osu.ppy.sh/mp/222
    Semifinals
    https://osu.ppy.sh/community/matches/333
    https://osu.ppy.sh/community/matches/111 duplicate"""
    links = extract_links_from_text(text)
    assert [(l.osu_match_id, normalize_round(l.round_raw)) for l in links] == [(111, "QF"), (222, "QF"), (333, "SF")]


def test_xlsx_hyperlinks_and_formulas(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws.append(["Round of 16"])
    ws.append(["Match 1", "Team A", "Team B", "MP Link"])
    ws["D2"].hyperlink = "https://osu.ppy.sh/community/matches/1001"
    ws.append(["Match 2", "Team C", "Team D", '=HYPERLINK("https://osu.ppy.sh/mp/1002","MP")'])
    ws.append(["Grand Finals", "Team A", "Team C", "https://osu.ppy.sh/community/matches/1003"])
    ws2 = wb.create_sheet("QF Results")
    ws2.append(["x", "https://osu.ppy.sh/community/matches/1004"])
    p = tmp_path / "s.xlsx"
    wb.save(p)
    kind, links = discover_links(str(p))
    got = {l.osu_match_id: normalize_round(l.round_raw) for l in links}
    assert kind == "xlsx"
    assert got == {1001: "RO16", 1002: "RO16", 1003: "GF", 1004: "QF"}


def test_google_sheet_url():
    u = "https://docs.google.com/spreadsheets/d/1AbC_d-9/edit#gid=0"
    assert google_sheet_export_url(u) == "https://docs.google.com/spreadsheets/d/1AbC_d-9/export?format=xlsx"


# ---------- parser ----------------------------------------------------------
def test_title_parse():
    assert parse_title("VRSO: (Team A) vs (Team B)") == ("VRSO", "Team A", "Team B")
    assert parse_title("OWC2025: (United States) vs. (Poland)") == ("OWC2025", "United States", "Poland")
    assert parse_title("random lobby") == (None, None, None)


def test_parser_handles_lazer_format():
    payload = {"match": {"id": 5, "name": "X: (a) vs (b)"}, "users": [{"id": 1, "username": "u"}], "events": [
        {"id": 1, "detail": {"type": "other"}, "game": {"id": 9, "beatmap_id": 77, "mods": [{"acronym": "HD"}],
         "scoring_type": "scorev2", "team_type": "head-to-head", "end_time": "x", "scores": [
             {"user_id": 1, "total_score": 900000, "legacy_total_score": 0, "accuracy": 0.98, "mods": [{"acronym": "HD"}],
              "statistics": {"great": 500, "ok": 3, "miss": 2}, "passed": True}]}}]}
    pm = parse_match(payload)
    s = pm.games[0].scores[0]
    assert pm.games[0].mods == ["HD"] and s.score == 900000 and s.count_miss == 2 and s.count_100 == 3


def test_infer_bucket():
    assert infer_bucket(["DT", "NF"], [["DT", "NF"], ["DT", "NF"]]) == "DT"
    assert infer_bucket(["NF"], [["HD", "NF"], ["HR", "NF"]]) == "FM"
    assert infer_bucket(["NF"], [["NF"], ["NF"]]) == "NM"
    assert infer_bucket(["NC", "HD"], [["NC", "HD"]] * 2) == "DT"


# ---------- end to end ------------------------------------------------------
@pytest.fixture()
def tourney(tmp_path):
    payloads, rounds = ft.build()
    conn = connect(tmp_path / "t.db")
    tid = repo.upsert_tournament(conn, "fake-2027", "Fake Open 2027", "FAKE", fmt="team")
    # source: a sheet-like text with section headers
    lines = []
    for rnd in ("Group Stage", "Semifinals", "Grand Finals"):
        lines.append(rnd)
        lines += [f"https://osu.ppy.sh/community/matches/{m}" for m, r in rounds.items() if r == rnd]
    lines.append("https://osu.ppy.sh/community/matches/424242")  # deleted lobby
    src = tmp_path / "links.txt"
    src.write_text("\n".join(lines))
    found, new = ingest.discover(conn, tid, str(src))
    assert found == new == len(payloads) + 1
    client = ft.FakeClient(payloads)
    counts = ingest.fetch(conn, tid, client, progress=lambda *_: None)
    assert counts == {"imported": len(payloads), "failed": 0, "not_found": 1}
    return conn, tid, client, payloads


def test_fetch_is_resumable_and_cached(tourney):
    conn, tid, client, payloads = tourney
    calls = client.calls
    assert ingest.fetch(conn, tid, client, progress=lambda *_: None)["imported"] == 0
    assert client.calls == calls
    assert ingest.reparse(conn, tid) == len(payloads)
    assert client.calls == calls  # reparse uses cache only


def test_mappool_marks_warmups(tourney):
    conn, tid, *_ = tourney
    rows = [{"beatmap_id": str(b), "slot": s} for b, s, _ in ft.POOL]
    repo.load_mappool(conn, tid, rows)
    ingest.finalize(conn, tid)
    warm = conn.execute("SELECT COUNT(*) FROM match_games WHERE beatmap_id = ? AND is_warmup = 1", (ft.WARMUP_MAP,)).fetchone()[0]
    games_with_warmup = conn.execute("SELECT COUNT(*) FROM match_games WHERE beatmap_id = ?", (ft.WARMUP_MAP,)).fetchone()[0]
    assert warm == games_with_warmup > 0
    fm = conn.execute("SELECT DISTINCT mod_bucket FROM match_games WHERE pool_slot LIKE 'FM%'").fetchall()
    assert [r[0] for r in fm] == ["FM"]


def test_report_ranks_known_skill(tourney):
    conn, tid, *_ = tourney
    repo.load_mappool(conn, tid, [{"beatmap_id": str(b), "slot": s} for b, s, _ in ft.POOL])
    ingest.finalize(conn, tid)
    rep = build_report(conn, "fake-2027")
    order = [r["username"] for r in rep["rankings"]]
    assert order[0] == "Alpha"                      # highest base skill
    assert order[-1] == "Foxtrot"                   # lowest base skill
    awards = {a["key"]: a["winner"] and a["winner"]["username"] for a in rep["awards"]}
    assert awards["mvp"] == "Alpha"
    assert awards["best_hr"] == "Bravo"             # HR specialist
    assert awards["best_hd"] in ("Delta", "Alpha")  # HD specialist vs best player
    assert rep["summary"]["rounds"] == ["GS", "SF", "GF"]
    page = rep["players"]["alpha"]
    assert page["rank"] == 1 and page["maps_played"] > 0 and page["match_history"]
    assert {m["round"] for m in rep["matches"]} == {"GS", "SF", "GF"}
    assert all(len(m["sides"]) == 2 for m in rep["matches"])
    team_a = next(t for t in rep["teams"] if t["name"] == "Team A")
    assert {p["username"] for p in team_a["players"]} == {"Alpha", "Bravo"}


def test_small_sample_shrinkage(tourney):
    """A player with 1 great map must not outrank a strong player with many maps."""
    from scout.analytics.ratings import Slice
    a, b = Slice(), Slice()
    a.add(2.5, 1.0)
    for _ in range(20):
        b.add(1.2, 1.0)
    assert a.z_mean() > b.z_mean()
    assert a.z_adj(3.0) < b.z_adj(3.0)


def test_decorated_headers_and_bare_ids():
    """Sheets like Corsace's: 'SEMIFINALS / BEST OF 1 | ━ ━ ━ | STATISTICS' headers and ids pasted without a URL."""
    from scout.sources.extract import Cell, Table, extract_links
    rule = "━  " * 30
    rows = [
        [Cell("MAP"), Cell("MP LINK")],
        [Cell("SEMIFINALS / BEST OF 1"), Cell(rule), Cell("STATISTICS")],
        [Cell("17"), Cell("criller"), Cell("+"), Cell("113747385")],
        [Cell("18"), Cell("x"), Cell("113749253", link="https://osu.ppy.sh/community/matches/113749253")],
        [Cell("GRAND FINALS / BEST OF 1"), Cell(rule)],
        [Cell("29"), Cell("enri"), Cell("113947963")],
    ]
    got = {l.osu_match_id: normalize_round(l.round_raw) for l in extract_links([Table("schedules", rows)])}
    assert got == {113747385: "SF", 113749253: "SF", 113947963: "GF"}
    # without an MP LINK column a bare number is not treated as a match id
    assert extract_links([Table("t", [[Cell("17"), Cell("113747385")]])]) == []
